# -*- coding: utf-8 -*-
"""Gera SPGG para a nuvem por edição cirúrgica do xlsx original (Office 365)."""

from __future__ import annotations

import html
import io
import re
import shutil
import zipfile
from datetime import datetime, date
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd

import gerar_teste_atualizacao as g

ORIG = g.SPGG_FILE
OUT = g.FOLDER / "SPGG_PARA_ENVIAR_NUVEM.xlsx"
OUT_CHECAGEM = g.FOLDER / "CHECAGEM_atualizacao_SPGG.xlsx"
SHEET2 = "xl/worksheets/sheet2.xml"
SHEET3 = "xl/worksheets/sheet3.xml"

CELL_RE = re.compile(r'<c r="([A-Z]+)(\d+)"([^>/]*)(?:/>|>(.*?)</c>)', re.S)


def xml_text(value) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    text = str(value)
    return escape(text, {"'": "&apos;", '"': "&quot;"})


def excel_serial(ts):
    ts = g.parse_date(ts)
    if pd.isna(ts):
        return None
    return int((pd.Timestamp(ts).normalize() - pd.Timestamp("1899-12-30")).days)


def style_of(attrs: str) -> str:
    m = re.search(r'\ss="[^"]*"', attrs)
    return m.group(0) if m else ""


def cell_inline(col, row, value, attrs="") -> str:
    s = style_of(attrs)
    if value is None or (isinstance(value, float) and pd.isna(value)) or str(value).strip() == "":
        return f'<c r="{col}{row}"{s}/>'
    return f'<c r="{col}{row}"{s} t="inlineStr"><is><t xml:space="preserve">{xml_text(value)}</t></is></c>'


def cell_num(col, row, value, attrs="") -> str:
    s = style_of(attrs)
    if value is None or (isinstance(value, float) and pd.isna(value)) or value == "":
        return f'<c r="{col}{row}"{s}/>'
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return f'<c r="{col}{row}"{s}><v>{value}</v></c>'


def convert_cell_xml(col, row, attrs, inner):
    if inner is None:
        return None
    if col == "N" and "<f" in inner:
        vm = re.search(r"<v>(.*?)</v>", inner, re.S)
        cached = vm.group(1) if vm else ""
        s = style_of(attrs)
        f = f'IF(J{row}="CC",L{row}*1950,IF(J{row}="FG",(L{row}*1950)*60%,""))'
        v = f"<v>{cached}</v>" if cached != "" else ""
        return f'<c r="N{row}"{s}><f>{f}</f>{v}</c>'
    if "<f" not in inner:
        return None
    if (
        "Banco de Talentos" not in inner
        and "[1]" not in inner
        and "[2]" not in inner
        and "TB_R_RH" not in inner
    ):
        return None
    vm = re.search(r"<v>(.*?)</v>", inner, re.S)
    v = vm.group(1) if vm else ""
    s = style_of(attrs)
    if v == "":
        return f'<c r="{col}{row}"{s}/>'
    if re.fullmatch(r"-?\d+(?:\.\d+)?(?:E[+-]?\d+)?", v):
        return f'<c r="{col}{row}"{s}><v>{v}</v></c>'
    raw = html.unescape(v)
    return f'<c r="{col}{row}"{s} t="inlineStr"><is><t xml:space="preserve">{xml_text(raw)}</t></is></c>'


def rewrite_sheet(xml: str, overrides: dict) -> str:
    parts = []
    last = 0
    for m in CELL_RE.finditer(xml):
        col, row_s, attrs, inner = m.group(1), m.group(2), m.group(3), m.group(4)
        row = int(row_s)
        new = overrides.get((col, row))
        if new is None:
            new = convert_cell_xml(col, row_s, attrs, inner)
        if new is None:
            continue
        parts.append(xml[last : m.start()])
        parts.append(new)
        last = m.end()
    parts.append(xml[last:])
    return "".join(parts)


def existing_attrs(xml, col, row) -> str:
    m = re.search(rf'<c r="{col}{row}"([^>/]*)', xml)
    return m.group(1) if m else ""


def replace_cell(xml, col, row, new_xml):
    pat = re.compile(rf'<c r="{col}{row}"[^>/]*(?:/>|>.*?</c>)', re.S)
    if pat.search(xml):
        return pat.sub(lambda _m: new_xml, xml, count=1)
    # insert before end of row
    row_pat = re.compile(rf'(<row r="{row}"[^>]*>)(.*?)(</row>)', re.S)
    m = row_pat.search(xml)
    if not m:
        return xml
    return xml[: m.start(3)] + new_xml + xml[m.start(3) :]


def overrides_for_item(item, row, attrs_map) -> dict:
    def a(col):
        return attrs_map.get((col, row), "")

    out = {
        ("G", row): cell_num("G", row, item["id_rh"], a("G")),
        ("H", row): cell_inline("H", row, item["nome"], a("H")),
        ("Z", row): cell_num("Z", row, item["id_ant"], a("Z")),
        ("AA", row): cell_inline("AA", row, item["saiu"], a("AA")),
        ("AF", row): cell_inline("AF", row, item["tel"] if item["nome"] else None, a("AF")),
    }
    serial = excel_serial(item["data"])
    if serial is not None:
        out[("F", row)] = cell_num("F", row, serial, a("F"))
    if item["status"]:
        out[("E", row)] = cell_inline("E", row, item["status"], a("E"))
    if item["cc_fg"]:
        out[("J", row)] = cell_inline("J", row, item["cc_fg"], a("J"))
    if item["codigo"]:
        out[("K", row)] = cell_inline("K", row, item["codigo"], a("K"))
    # ENCARGO fica na VAGA_1; não atualizar em linha existente.
    if item["nome"]:
        out[("Q", row)] = cell_inline("Q", row, item["via"], a("Q"))
        out[("R", row)] = cell_inline("R", row, item["carg"], a("R"))
        out[("S", row)] = cell_inline("S", row, item["indic"], a("S"))
        out[("T", row)] = cell_inline("T", row, item["partido"], a("T"))
    if item["saiu"]:
        out[("AB", row)] = cell_inline("AB", row, item["via_ant"], a("AB"))
        out[("AC", row)] = cell_inline("AC", row, item["carg_ant"], a("AC"))
        out[("AD", row)] = cell_inline("AD", row, item["indic_ant"], a("AD"))
        out[("AE", row)] = cell_inline("AE", row, item["partido_ant"], a("AE"))
    if item.get("fator") is not None:
        out[("L", row)] = cell_num("L", row, item["fator"], a("L"))
    return out


def build_new_row(item, row):
    serial = excel_serial(item["data"]) or ""
    nome = item["nome"]
    enc = item["encargo"] or ""
    orgao = item["orgao"] or ""
    unidade = (str(item["unidade"]).replace("\n", " ") if item["unidade"] else "")
    chefia = (
        "CHEFIA"
        if enc and re.search(r"diretor|chefe de gabinete", str(enc), re.I)
        else "ASS."
    )
    cells = [
        cell_num("A", row, item["vaga"], ' s="127"'),
        cell_inline("B", row, orgao, ' s="138"'),
        cell_inline("C", row, "DIRETA", ' s="138"'),
        cell_inline("D", row, chefia, ' s="145"'),
        cell_inline("E", row, item["status"] or "OCUPADO", ' s="145"'),
        cell_num("F", row, serial if serial != "" else None, ' s="148"'),
        cell_num("G", row, item["id_rh"], ' s="145"'),
        cell_inline("H", row, nome, ' s="141"'),
        f'<c r="I{row}" s="145"/>',
        cell_inline("J", row, item["cc_fg"], ' s="145"'),
        cell_inline("K", row, item["codigo"], ' s="138"'),
        cell_num("L", row, item.get("fator"), ' s="138"'),
        f'<c r="M{row}" s="142"><f>RIGHT(K{row}, 2)</f></c>',
        f'<c r="N{row}" s="143"><f>IF(J{row}="CC",L{row}*1950,IF(J{row}="FG",(L{row}*1950)*60%,""))</f></c>',
        cell_inline("O", row, enc, ' s="139"'),
        cell_inline("P", row, unidade, ' s="139"'),
        cell_inline("Q", row, item["via"] if nome else "-", ' s="321"'),
        cell_inline("R", row, item["carg"] if nome else "-", ' s="321"'),
        cell_inline("S", row, item["indic"] if nome else "-", ' s="321"'),
        cell_inline("T", row, item["partido"] if nome else "-", ' s="321"'),
        f'<c r="U{row}" s="321"><f>T{row}</f></c>',
        f'<c r="V{row}" s="236"/>',
        f'<c r="W{row}" s="236"/>',
        f'<c r="X{row}" s="127"/>',
        cell_num("Y", row, item["vaga"], ' s="127"'),
        cell_num("Z", row, item["id_ant"], ' s="256"'),
        cell_inline("AA", row, item["saiu"], ' s="251"'),
        cell_inline("AB", row, item["via_ant"] if item["saiu"] else "-", ' s="251"'),
        cell_inline("AC", row, item["carg_ant"] if item["saiu"] else "-", ' s="252"'),
        cell_inline("AD", row, item["indic_ant"] if item["saiu"] else "-", ' s="252"'),
        cell_inline("AE", row, item["partido_ant"] if item["saiu"] else "-", ' s="252"'),
        cell_inline("AF", row, item["tel"] if nome else None, ' s="252"'),
    ]
    inner = "".join(cells)
    return f'<row r="{row}" spans="1:32">{inner}</row>'


def map_vaga_rows(xml: str) -> dict:
    mapping = {}
    for m in re.finditer(r'<c r="A(\d+)"[^>]*><v>(\d+)</v></c>', xml):
        mapping[int(m.group(2))] = int(m.group(1))
    return mapping


def last_row_number(xml: str) -> int:
    nums = [int(n) for n in re.findall(r'<row r="(\d+)"', xml)]
    return max(nums) if nums else 1


def patch_workbook_xml(xml: str, last_row: int) -> str:
    xml = re.sub(r"<externalReferences>.*?</externalReferences>", "", xml, flags=re.S)
    xml = re.sub(
        r'<definedName name="Padrao">.*?</definedName>',
        "",
        xml,
        flags=re.S,
    )
    xml = re.sub(
        r"('TUDO POR VAGA'!\$A\$1:\$AF\$)\d+",
        rf"\g<1>{last_row}",
        xml,
    )
    if "fullCalcOnLoad" not in xml:
        xml = xml.replace("<calcPr", '<calcPr fullCalcOnLoad="1"')
    return xml


def patch_rels(xml: str) -> str:
    xml = re.sub(
        r'<Relationship[^>]+Target="externalLinks/externalLink\d+\.xml"[^/]*/>',
        "",
        xml,
    )
    xml = re.sub(
        r'<Relationship[^>]+Target="calcChain\.xml"[^/]*/>',
        "",
        xml,
    )
    return xml


def patch_content_types(xml: str) -> str:
    xml = re.sub(r'<Override[^>]+calcChain\.xml"[^/]*/>', "", xml)
    xml = re.sub(r'<Override[^>]+externalLink\d+\.xml"[^/]*/>', "", xml)
    return xml


def _style_sheet(ws, n_cols, fill_hex):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    for col in range(1, n_cols + 1):
        cell = ws.cell(1, col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        longest = 12
        for cell in col[:60]:
            if cell.value:
                longest = max(longest, min(len(str(cell.value)) + 2, 42))
        ws.column_dimensions[letter].width = longest


def _norm_cell(v):
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return str(v).strip()


def snapshot_checagem_anterior():
    if not OUT_CHECAGEM.exists():
        return pd.DataFrame(), pd.DataFrame()
    try:
        novas = pd.read_excel(OUT_CHECAGEM, sheet_name="Vagas_novas", dtype=str)
    except Exception:
        novas = pd.DataFrame()
    try:
        movs = pd.read_excel(OUT_CHECAGEM, sheet_name="Movimentacoes", dtype=str)
    except Exception:
        movs = pd.DataFrame()
    return novas.fillna(""), movs.fillna("")


def diff_checagem(prev_n, prev_m, df_n, df_m):
    rows = []

    def key_n(df):
        if df is None or df.empty or "VAGA_1" not in df.columns:
            return {}
        out = {}
        for _, r in df.iterrows():
            k = _norm_cell(r.get("VAGA_1"))
            if k:
                out[k] = r
        return out

    pn, pm = key_n(prev_n), key_n(prev_m)
    nn, nm = key_n(df_n), key_n(df_m)

    def linha(aba, tipo, vaga, antes, depois):
        def get(rec, *cols):
            if rec is None:
                return ""
            for c in cols:
                if c in rec.index:
                    return _norm_cell(rec.get(c))
            return ""

        return {
            "Aba": aba,
            "Tipo": tipo,
            "VAGA_1": vaga,
            "Antes — nome": get(antes, "Nome", "Nome que entra"),
            "Antes — ato": get(antes, "Ato"),
            "Antes — data": get(antes, "DATA DOE"),
            "Antes — quem sai": get(antes, "Nome que sai"),
            "Agora — nome": get(depois, "Nome", "Nome que entra"),
            "Agora — ato": get(depois, "Ato"),
            "Agora — data": get(depois, "DATA DOE"),
            "Agora — quem sai": get(depois, "Nome que sai"),
        }

    for v in sorted(set(nn) - set(pn), key=lambda x: int(x) if x.isdigit() else 0):
        rows.append(linha("Vagas_novas", "INCLUIDA", v, None, nn[v]))
    for v in sorted(set(pn) - set(nn), key=lambda x: int(x) if x.isdigit() else 0):
        rows.append(linha("Vagas_novas", "REMOVIDA", v, pn[v], None))
    for v in sorted(set(nn) & set(pn), key=lambda x: int(x) if x.isdigit() else 0):
        a, b = pn[v], nn[v]
        campos = [
            c
            for c in ("Nome", "Nome que entra", "Nome que sai", "Ato", "DATA DOE", "STATUS", "CC/FG", "CÓDIGO")
            if c in a.index or c in b.index
        ]
        mudou = False
        for c in campos:
            if _norm_cell(a.get(c) if c in a.index else "") != _norm_cell(b.get(c) if c in b.index else ""):
                mudou = True
                break
        if mudou:
            rows.append(linha("Vagas_novas", "ALTERADA", v, a, b))

    for v in sorted(set(nm) - set(pm), key=lambda x: int(x) if x.isdigit() else 0):
        rows.append(linha("Movimentacoes", "INCLUIDA", v, None, nm[v]))
    for v in sorted(set(pm) - set(nm), key=lambda x: int(x) if x.isdigit() else 0):
        rows.append(linha("Movimentacoes", "REMOVIDA", v, pm[v], None))
    for v in sorted(set(nm) & set(pm), key=lambda x: int(x) if x.isdigit() else 0):
        a, b = pm[v], nm[v]
        campos = ["Nome que entra", "Nome que sai", "Ato", "DATA DOE", "STATUS", "CC/FG", "CÓDIGO"]
        mudou = False
        for c in campos:
            if _norm_cell(a.get(c) if c in a.index else "") != _norm_cell(b.get(c) if c in b.index else ""):
                mudou = True
                break
        if mudou:
            rows.append(linha("Movimentacoes", "ALTERADA", v, a, b))

    return pd.DataFrame(rows)


def write_checagem(carga, n_upd, n_new, nota=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils.dataframe import dataframe_to_rows

    prev_n, prev_m = snapshot_checagem_anterior()

    novas = [x for x in carga if x.get("nova")]
    movs = [x for x in carga if not x.get("nova")]

    def vaga_txt(v):
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return ""
        return str(int(v))

    df_novas = pd.DataFrame(
        [
            {
                "VAGA_1": vaga_txt(x["vaga"]),
                "Nome": x["nome"],
                "ID RH": x["id_rh"],
                "Nome que sai": x["saiu"],
                "ID_RH ANT.": x["id_ant"],
                "DATA DOE": g.fmt_date(x["data"]),
                "Ato": x["ato"],
                "STATUS": x["status"],
                "CC/FG": x["cc_fg"],
                "CÓDIGO": x["codigo"],
                "ENCARGO (da vaga nova)": x["encargo"],
                "Órgão": x["orgao"],
            }
            for x in novas
        ]
    ).sort_values("VAGA_1", key=lambda s: s.map(lambda x: int(x) if str(x).isdigit() else 0)) if novas else pd.DataFrame()

    df_movs = pd.DataFrame(
        [
            {
                "VAGA_1": vaga_txt(x["vaga"]),
                "Nome que entra": x["nome"],
                "ID RH": x["id_rh"],
                "Nome que sai": x["saiu"],
                "ID_RH ANT.": x["id_ant"],
                "DATA DOE": g.fmt_date(x["data"]),
                "Ato": x["ato"],
                "STATUS": x["status"],
                "CC/FG": x["cc_fg"],
                "CÓDIGO": x["codigo"],
                "Órgão SPGG": x["orgao"],
            }
            for x in movs
        ]
    ).sort_values("VAGA_1", key=lambda s: s.map(lambda x: int(x) if str(x).isdigit() else 0)) if movs else pd.DataFrame()

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "LEIA-ME"
    ws0["A1"] = "Checagem da atualização SPGG × DOE"
    ws0["A1"].font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
    ws0.merge_cells("A1:F1")
    lines = [
        "",
        "Planilha só para conferir. Não envie este arquivo para a nuvem.",
        "O arquivo da nuvem é: SPGG_PARA_ENVIAR_NUVEM.xlsx",
        nota or "",
        "",
        "ABAS",
        "1. Vagas_novas — VAGA_1 que não existia na SPGG (primeiro provimento / função nova).",
        "2. Movimentacoes — vagas que já existiam e tiveram troca de titular, data, CC/FG ou código.",
        "3. O_que_mudou — diferença desta checagem em relação à anterior (incluída / removida / alterada).",
        "",
        "Uma fonte por geração: somente FG_Cargos_Comissao.xlsx (não mistura com o Atualizado).",
        "Premissa: se a movimentação está nesse extrato, ela aconteceu.",
        "Titular = NOMEIA/DESIGNA/COLOCA/ATRIBUI/EXONERA/DISPENSA/CESSA.",
        "RETIFICA de tipo FG/CC só altera CC/FG e CÓDIGO; não nomeia ninguém.",
        "SEM EFEITO só tira titular em posse/desistência/publicação indevida.",
        "APOSTILA é correção textual — não conta como movimentação de titular.",
        "ENCARGO nas vagas já existentes não muda (é da VAGA_1, não do nome).",
        "ID RH em branco = pessoa que não estava na SPGG.",
        "VAGA_1 está em texto completo (sem notação científica).",
        "",
        f"Vagas novas: {len(novas)}",
        f"Movimentações em vagas existentes: {len(movs)}",
        f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
    ]
    for i, line in enumerate(lines, start=2):
        ws0[f"A{i}"] = line
        ws0[f"A{i}"].font = Font(
            name="Calibri",
            size=12,
            bold=line in {"ABAS"},
            color="1F4E79" if line in {"ABAS"} else "000000",
        )
    ws0.column_dimensions["A"].width = 120

    def dump(title, df, color):
        ws = wb.create_sheet(title)
        if df is None or df.empty:
            ws["A1"] = "Nenhum registro"
            return
        view = df.copy()
        for col in view.columns:
            view[col] = view[col].apply(lambda v: "" if pd.isna(v) else v)
        headers = list(view.columns)
        id_cols = {i + 1 for i, h in enumerate(headers) if h in {"ID RH", "ID_RH ANT."}}
        vaga_cols = {i + 1 for i, h in enumerate(headers) if h == "VAGA_1"}
        for r_idx, row in enumerate(dataframe_to_rows(view, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(r_idx, c_idx, value)
                cell.font = Font(name="Calibri", size=10)
                if r_idx > 1 and c_idx in vaga_cols:
                    cell.number_format = "@"
                    if value not in ("", None):
                        cell.value = str(value)
                if r_idx > 1 and c_idx in id_cols and value not in ("", None):
                    try:
                        cell.value = int(float(value))
                        cell.number_format = "0"
                    except (TypeError, ValueError):
                        pass
        _style_sheet(ws, view.shape[1], color)
        if "A" in ws.column_dimensions:
            ws.column_dimensions["A"].width = 16

    dump("Vagas_novas", df_novas, "548235")
    dump("Movimentacoes", df_movs, "C65911")
    df_alt = diff_checagem(prev_n, prev_m, df_novas, df_movs)
    dump("O_que_mudou", df_alt, "1F4E79")
    if "LEIA-ME" in wb.sheetnames:
        ws0 = wb["LEIA-ME"]
        r = ws0.max_row + 2
        ws0[f"A{r}"] = (
            f"Alterações vs checagem anterior: {0 if df_alt is None or df_alt.empty else len(df_alt)} "
            f"(incluídas/removidas/alteradas)."
        )
        ws0[f"A{r}"].font = Font(name="Calibri", size=12, bold=True, color="1F4E79")
    wb.save(OUT_CHECAGEM)
    print("OK checagem", OUT_CHECAGEM, "novas", len(novas), "movs", len(movs), "mudancas", 0 if df_alt is None or df_alt.empty else len(df_alt))


def main():
    print("Cruzamento...")
    doe = g.load_all_doe()
    selected, discarded, sem_vaga, so_lotacao = g.select_last_acts(doe)
    spgg = pd.read_excel(ORIG, sheet_name="TUDO POR VAGA")
    update_map, audit, no_match, _idx = g.build_updates(selected, spgg)
    packs = g.person_pack(spgg)
    carga = g.build_carga(audit, update_map, packs)

    print("Copiando original...")
    if OUT.exists():
        OUT.unlink()
    shutil.copy2(ORIG, OUT)

    with zipfile.ZipFile(ORIG, "r") as zin:
        sheet2 = zin.read(SHEET2).decode("utf-8")
        sheet3 = zin.read(SHEET3).decode("utf-8")
        wbxml = zin.read("xl/workbook.xml").decode("utf-8")
        rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        ctypes = zin.read("[Content_Types].xml").decode("utf-8")

    print("Mapeando células...")
    attrs_map = {}
    row_of = {}
    for m in CELL_RE.finditer(sheet2):
        col, row_s, attrs, inner = m.group(1), m.group(2), m.group(3), m.group(4)
        row = int(row_s)
        attrs_map[(col, row)] = attrs
        if col == "A" and inner:
            vm = re.search(r"<v>(\d+)</v>", inner)
            if vm:
                row_of[int(vm.group(1))] = row
    print("vagas mapeadas", len(row_of))

    overrides = {}
    n_upd = 0
    novas = []
    for item in carga:
        r = row_of.get(item["vaga"])
        if r:
            overrides.update(overrides_for_item(item, r, attrs_map))
            n_upd += 1
        else:
            novas.append(item)

    print("Reescrevendo TUDO POR VAGA (sem PROCV externo)...")
    sheet2 = rewrite_sheet(sheet2, overrides)
    print("Reescrevendo INDIRETA...")
    sheet3 = rewrite_sheet(sheet3, {})

    last = last_row_number(sheet2)
    new_rows_xml = []
    n_new = 0
    for item in novas:
        if item["vaga"] in row_of:
            continue
        last += 1
        new_rows_xml.append(build_new_row(item, last))
        n_new += 1
    if new_rows_xml:
        sheet2 = sheet2.replace("</sheetData>", "".join(new_rows_xml) + "</sheetData>")
        sheet2 = re.sub(r'<dimension ref="A1:AF\d+"/>', f'<dimension ref="A1:AF{last}"/>', sheet2)
        sheet2 = re.sub(r'(<autoFilter ref="A1:AF)\d+"', rf'\g<1>{last}"', sheet2)

    wbxml = patch_workbook_xml(wbxml, last)
    rels = patch_rels(rels)
    ctypes = patch_content_types(ctypes)

    skip = {
        SHEET2,
        SHEET3,
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "[Content_Types].xml",
        "xl/calcChain.xml",
        "xl/externalLinks/externalLink1.xml",
        "xl/externalLinks/externalLink2.xml",
        "xl/externalLinks/_rels/externalLink1.xml.rels",
        "xl/externalLinks/_rels/externalLink2.xml.rels",
    }

    print("Regravando xlsx...")
    tmp = OUT.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(ORIG, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            if item.filename in skip:
                continue
            zout.writestr(item, zin.read(item.filename))
        zout.writestr(SHEET2, sheet2.encode("utf-8"))
        zout.writestr(SHEET3, sheet3.encode("utf-8"))
        zout.writestr("xl/workbook.xml", wbxml.encode("utf-8"))
        zout.writestr("xl/_rels/workbook.xml.rels", rels.encode("utf-8"))
        zout.writestr("[Content_Types].xml", ctypes.encode("utf-8"))

    tmp.replace(OUT)
    print("OK", OUT)
    print("atualizadas", n_upd, "novas", n_new, "ultima linha", last)
    win = doe[(doe["data_p"] >= "2026-05-11") & (doe["data_p"] < "2027-01-01")]
    dt_min = pd.Timestamp(win["data_p"].min()).strftime("%d/%m/%Y") if len(win) else "?"
    dt_max = pd.Timestamp(win["data_p"].max()).strftime("%d/%m/%Y") if len(win) else "?"
    fonte = g.resolve_doe_fonte()
    nota = (
        f"Período completo: {dt_min} até {dt_max}. "
        f"Fonte única: {fonte}."
    )
    write_checagem(carga, n_upd, n_new, nota=nota)


def patch_xlsx(base: Path, carga, skip_later: bool = True):
    spgg_now = pd.read_excel(base, sheet_name="TUDO POR VAGA")
    current_dates = {}
    for _, r in spgg_now.iterrows():
        k = g.vaga_int(r.get("VAGA_1"))
        if k is not None:
            current_dates[k] = g.parse_date(r.get("DATA DOE"))

    applied = []
    skipped_later = []
    for item in carga:
        cur = current_dates.get(item["vaga"])
        if (
            skip_later
            and pd.notna(cur)
            and pd.notna(item["data"])
            and pd.Timestamp(cur).normalize() > pd.Timestamp(item["data"]).normalize()
        ):
            skipped_later.append(item)
            continue
        applied.append(item)

    print("Copiando base", base.name, "...")
    tmp_src = base.with_suffix(".bak_src.xlsx")
    shutil.copy2(base, tmp_src)

    with zipfile.ZipFile(tmp_src, "r") as zin:
        sheet2 = zin.read(SHEET2).decode("utf-8")
        sheet3 = zin.read(SHEET3).decode("utf-8")
        wbxml = zin.read("xl/workbook.xml").decode("utf-8")
        rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        ctypes = zin.read("[Content_Types].xml").decode("utf-8")

    attrs_map = {}
    row_of = {}
    for m in CELL_RE.finditer(sheet2):
        col, row_s, attrs, inner = m.group(1), m.group(2), m.group(3), m.group(4)
        row = int(row_s)
        attrs_map[(col, row)] = attrs
        if col == "A" and inner:
            vm = re.search(r"<v>(\d+)</v>", inner)
            if vm:
                row_of[int(vm.group(1))] = row

    overrides = {}
    n_upd = 0
    novas = []
    for item in applied:
        r = row_of.get(item["vaga"])
        if r:
            overrides.update(overrides_for_item(item, r, attrs_map))
            n_upd += 1
        else:
            novas.append(item)

    sheet2 = rewrite_sheet(sheet2, overrides)
    sheet3 = rewrite_sheet(sheet3, {})

    last = last_row_number(sheet2)
    new_rows_xml = []
    n_new = 0
    for item in novas:
        if item["vaga"] in row_of:
            continue
        last += 1
        new_rows_xml.append(build_new_row(item, last))
        n_new += 1
    if new_rows_xml:
        sheet2 = sheet2.replace("</sheetData>", "".join(new_rows_xml) + "</sheetData>")
        sheet2 = re.sub(r'<dimension ref="A1:AF\d+"/>', f'<dimension ref="A1:AF{last}"/>', sheet2)
        sheet2 = re.sub(r'(<autoFilter ref="A1:AF)\d+"', rf'\g<1>{last}"', sheet2)

    wbxml = patch_workbook_xml(wbxml, last)
    rels = patch_rels(rels)
    ctypes = patch_content_types(ctypes)

    skip = {
        SHEET2,
        SHEET3,
        "xl/workbook.xml",
        "xl/_rels/workbook.xml.rels",
        "[Content_Types].xml",
        "xl/calcChain.xml",
        "xl/externalLinks/externalLink1.xml",
        "xl/externalLinks/externalLink2.xml",
        "xl/externalLinks/_rels/externalLink1.xml.rels",
        "xl/externalLinks/_rels/externalLink2.xml.rels",
    }

    tmp = OUT.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(tmp_src, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename in skip:
                continue
            zout.writestr(info, zin.read(info.filename))
        zout.writestr(SHEET2, sheet2.encode("utf-8"))
        zout.writestr(SHEET3, sheet3.encode("utf-8"))
        zout.writestr("xl/workbook.xml", wbxml.encode("utf-8"))
        zout.writestr("xl/_rels/workbook.xml.rels", rels.encode("utf-8"))
        zout.writestr("[Content_Types].xml", ctypes.encode("utf-8"))
    tmp.replace(OUT)
    tmp_src.unlink(missing_ok=True)
    print("OK", OUT, "atualizadas", n_upd, "novas", n_new, "ignoradas (data posterior na SPGG)", len(skipped_later))
    return applied, skipped_later, n_upd, n_new


def _carga_from_checagem_row(r, nova: bool) -> dict:
    nome_col = "Nome" if nova else "Nome que entra"
    orgao_col = "Órgão" if nova else "Órgão SPGG"
    enc_col = "ENCARGO (da vaga nova)" if nova else None
    return {
        "vaga": int(g.vaga_int(r.get("VAGA_1"))),
        "nova": nova,
        "nome": r.get(nome_col) if nome_col in r.index else r.get("Nome que entra"),
        "saiu": r.get("Nome que sai"),
        "id_rh": g.as_id(r.get("ID RH")),
        "id_ant": g.as_id(r.get("ID_RH ANT.")),
        "data": g.parse_date(r.get("DATA DOE")),
        "status": r.get("STATUS"),
        "cc_fg": r.get("CC/FG"),
        "codigo": r.get("CÓDIGO"),
        "encargo": r.get(enc_col) if enc_col and enc_col in r.index else None,
        "orgao": r.get(orgao_col) if orgao_col in r.index else r.get("Órgão"),
        "ato": r.get("Ato"),
    }


def incremental_lote(doe_path: Path):
    print("Lote incremental:", doe_path.name)
    doe = g.load_doe(doe_path, linha_offset=200000)
    selected, discarded, sem_vaga, so_lotacao = g.select_last_acts(doe)
    base = OUT if OUT.exists() else ORIG
    spgg = pd.read_excel(base, sheet_name="TUDO POR VAGA")
    update_map, audit, no_match, _idx = g.build_updates(selected, spgg)
    packs = g.person_pack(spgg)
    carga = g.build_carga(audit, update_map, packs)
    print("atos no lote (último por vaga)", len(carga))
    prev_n, prev_m = snapshot_checagem_anterior()
    applied, skipped, n_upd, n_new = patch_xlsx(base, carga, skip_later=True)
    merged = {}
    if prev_n is not None and not prev_n.empty:
        for _, r in prev_n.iterrows():
            k = g.vaga_int(r.get("VAGA_1"))
            if k:
                merged[int(k)] = _carga_from_checagem_row(r, True)
    if prev_m is not None and not prev_m.empty:
        for _, r in prev_m.iterrows():
            k = g.vaga_int(r.get("VAGA_1"))
            if k:
                merged[int(k)] = _carga_from_checagem_row(r, False)
    for item in applied:
        merged[int(item["vaga"])] = item
    nota = (
        f"Lote incremental: {doe_path.name}. "
        f"Aplicadas {n_upd} vagas existentes e {n_new} novas. "
        f"Ignoradas (data posterior na SPGG): {len(skipped)}. "
        "Checagem = histórico anterior + este lote."
    )
    write_checagem(list(merged.values()), n_upd, n_new, nota=nota)


if __name__ == "__main__":
    main()
