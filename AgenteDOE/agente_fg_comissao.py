#!/usr/bin/env python3
"""
Agente FG / Cargos em Comissão - Versão Interativa (Corrigida)
- Seleção de múltiplos arquivos
- Cancelar com 0 / sair / cancelar
- Sem duplicação de PDFs
- Extração de quadros Nome / Código / Vaga / Encargo / Titular Anterior
"""

import argparse
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional, Tuple
from collections import defaultdict

PASTA_AGENTE = Path(__file__).resolve().parent

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


COLUNAS = [
    "Nome", "Assunto", "Indicação nova", "Data DOE", "A contar de", "Indicação",
    "Partido dono da vaga", "Deputado", "Vaga Anterior",
    "Proa", "Vaga n°", "Função", "Órgão / Secretaria"
]

ASSUNTO_PADRAO = "Cargos em Comissão e Funções Gratificadas"
ASSUNTO_FG_RE = re.compile(
    r"cargos?\s+em\s+comiss[aã]o\s+e\s+fun[cç][oõ]es?\s+gratificadas?",
    re.IGNORECASE,
)
ASSUNTO_IGNORAR_RE = re.compile(
    r"substitui[cç][aã]o|posse|afastamento|licen[cç]a|contrato\s+tempor|"
    r"abono|concurso|portaria|convoca[cç][aã]o|carga\s+hor[aá]ria|"
    r"inclus[aã]o|atos\s+administrativos",
    re.IGNORECASE,
)
SEM_EFEITO_RE = re.compile(r"\b(TORNA|DECLARA)\s+SEM\s+EFEITO", re.IGNORECASE)

# CESSA/DISPENSA/EXONERA desocupam a vaga; ATRIBUI/DESIGNA/NOMEIA ocupam.
INDICACOES_ORDEM = [
    "RETIFICA", "ALTERA",
    "CESSA", "EXONERA", "DISPENSA",
    "NOMEIA", "DESIGNA", "COLOCA", "ATRIBUI", "RELOTA",
]
INDICACOES_NOVA = {"NOMEIA", "DESIGNA", "ATRIBUI"}
INDICACOES_DESOCUPA = {"EXONERA", "DISPENSA", "CESSA"}
INDICACOES_ENTRADA_VAGA = {"NOMEIA", "DESIGNA", "ATRIBUI"}

CODIGO_QUADRO_RE = re.compile(r"\b(FGT|FGS|CCT|CCS)\s*[-–/]?\s*(\d{1,2})\b", re.IGNORECASE)


def padronizar_codigo_funcao(texto: str) -> str:
    """Cargo Comissionado Superior 11 → CCS 11 (idem CCT, FGS, FGT)."""
    if not texto:
        return ""
    t = re.sub(r"\s+", " ", str(texto)).strip()
    pares = [
        (r"cargo\s+comissionado\s+superior", "CCS"),
        (r"cargo\s+comiss(?:ionado|\.)?\s*sup(?:erior|\.)?", "CCS"),
        (r"cargo\s+comissionado\s+transversal", "CCT"),
        (r"cargo\s+comiss(?:ionado|\.)?\s*transv(?:ersal|\.)?", "CCT"),
        (r"fun[cç][aã]o\s+gratif(?:icada|\.)?\s*(?:superior|sup\.?)", "FGS"),
        (r"fun[cç][aã]o\s+gratif(?:icada|\.)?\s*(?:transversal|transv\.?)", "FGT"),
        (r"\bccs\b", "CCS"),
        (r"\bcct\b", "CCT"),
        (r"\bfgs\b", "FGS"),
        (r"\bfgt\b", "FGT"),
    ]
    for pat, sigla in pares:
        m = re.search(pat + r"\s*[-–/]?\s*(\d{1,2})\b", t, re.IGNORECASE)
        if m:
            return f"{sigla} {int(m.group(1)):02d}"
    m = re.search(r"\b(CCS|CCT|FGS|FGT)\s*[-–/]?\s*(\d{1,2})\b", t, re.IGNORECASE)
    if m:
        return f"{m.group(1).upper()} {int(m.group(2)):02d}"
    return ""
VAGA_QUADRO_RE = re.compile(r"\b(\d{9,14})\b")
ID_FUNCIONAL_RE = re.compile(r"\b\d{5,8}\s*/\s*\d{1,2}\b")


def listar_pdfs() -> List[Path]:
    pasta = PASTA_AGENTE
    pdfs = sorted(set(pasta.glob("*.pdf")) | set(pasta.glob("*.PDF")))
    return pdfs


def resolver_pdfs_cli(nomes: List[str]) -> List[Path]:
    selecionados = []
    for nome in nomes:
        caminho = Path(nome)
        if not caminho.is_absolute():
            caminho = PASTA_AGENTE / caminho
        if not caminho.exists():
            print(f"  [ERRO] PDF não encontrado: {nome}")
            sys.exit(1)
        selecionados.append(caminho.resolve())
    return selecionados


def selecionar_pdfs(pdfs: List[Path]) -> List[Path]:
    if not pdfs:
        print("Nenhum arquivo PDF encontrado na pasta atual.")
        sys.exit(0)

    print("\nArquivos PDF encontrados:\n")
    for i, pdf in enumerate(pdfs, 1):
        print(f"  {i:2d}. {pdf.name}")

    print("\nDigite os números dos arquivos que deseja processar")
    print("Pode separar por ESPAÇO ou VÍRGULA")
    print("Exemplos:  1 3 5   |   1,3,5   |   all")
    print("Para cancelar digite: 0  ou  sair  ou  cancelar\n")

    escolha = input("Sua escolha: ").strip().lower()

    if escolha in ["0", "sair", "cancelar", "voltar", "exit", "quit"]:
        print("\nOperação cancelada pelo usuário.")
        sys.exit(0)

    if escolha == "all":
        return pdfs

    escolha = escolha.replace(",", " ")
    partes = escolha.split()

    selecionados = []
    try:
        for parte in partes:
            num = int(parte)
            if 1 <= num <= len(pdfs):
                selecionados.append(pdfs[num - 1])
            else:
                print(f"  [AVISO] Número {num} inválido, ignorado.")
    except ValueError:
        print("Entrada inválida. Use apenas números ou 'all'.")
        sys.exit(1)

    if not selecionados:
        print("Nenhum arquivo válido selecionado.")
        sys.exit(0)

    return selecionados


def extrair_texto_com_paginas(caminho_pdf: str) -> List[Dict]:
    resultado = []
    try:
        with pdfplumber.open(caminho_pdf) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                texto = page.extract_text() or ""
                resultado.append({"pagina": i, "texto": texto})
    except Exception as e:
        print(f"  [ERRO] Falha ao extrair {caminho_pdf}: {e}")
        return []
    return resultado


def extrair_sumario(paginas: List[Dict]) -> List[Tuple[str, int]]:
    secoes = []
    padrao = r"(.+?)\.{2,}\s*(\d+)"
    for pagina in paginas[:6]:
        for linha in pagina["texto"].split("\n"):
            match = re.search(padrao, linha)
            if match:
                nome = match.group(1).strip()
                pagina_num = int(match.group(2))
                if any(x in nome.upper() for x in ["GOVERNO", "GABINETE", "SECRETARIA", "ATOS DO"]):
                    secoes.append((nome.strip(), pagina_num))
    vistos = set()
    return [s for s in secoes if not (s[0] in vistos or vistos.add(s[0]))]


def extrair_data_doe(texto: str) -> str:
    m = re.search(
        r"(\d{1,2})\s+de\s+(janeiro|fevereiro|março|abril|maio|junho|julho|agosto|setembro|outubro|novembro|dezembro)\s+de\s+(\d{4})",
        texto[:4000], re.IGNORECASE
    )
    if m:
        meses = {
            "janeiro": "01", "fevereiro": "02", "março": "03", "abril": "04",
            "maio": "05", "junho": "06", "julho": "07", "agosto": "08",
            "setembro": "09", "outubro": "10", "novembro": "11", "dezembro": "12"
        }
        return f"{m.group(1).zfill(2)}/{meses[m.group(2).lower()]}/{m.group(3)}"
    return datetime.now().strftime("%d/%m/%Y")


def extrair_nome(texto: str) -> str:
    m = re.search(
        r"Nome:\s*([A-ZÁÀÂÃÄÉÈÊËÍÌÎÏÓÒÔÕÖÚÙÛÜÇÑ][A-Za-záàâãäéèêëíìîïóòôõöúùûüçñ'\s-]+?)"
        r"(?:,|\n|Identificação|Lota[cç][aã]o|$)",
        texto,
        re.IGNORECASE,
    )
    if m:
        return re.sub(r"\s+", " ", m.group(1).strip())
    return ""


def limpar_assunto(texto: str) -> str:
    if not texto:
        return ""
    m = re.search(r"Assunto:\s*(.+)", texto, re.IGNORECASE | re.DOTALL)
    bruto = m.group(1) if m else texto
    bruto = re.split(
        r"\s+(?:Processo|Nome|Protocolo|Identifica[cç][aã]o|Tipo\s+V[ií]nculo|Cargo/Fun[cç][aã]o|Lota[cç][aã]o)\s*:",
        bruto,
        maxsplit=1,
        flags=re.IGNORECASE,
    )[0]
    bruto = re.split(r"\s+O\s+GOVERNADOR\b", bruto, maxsplit=1, flags=re.IGNORECASE)[0]
    return re.sub(r"\s+", " ", bruto).strip(" .;,-")


def eh_assunto_fg_comissao(texto: str) -> bool:
    assunto = limpar_assunto(texto)
    if not assunto:
        return False
    if ASSUNTO_IGNORAR_RE.search(assunto):
        return False
    return bool(ASSUNTO_FG_RE.search(assunto))


def extrair_sem_efeito(texto: str) -> str:
    m = SEM_EFEITO_RE.search(texto or "")
    if not m:
        return ""
    return f"{m.group(1).upper()} SEM EFEITO"


def tem_codigo_vaga(texto: str) -> bool:
    return bool(extrair_vaga_num(texto))


def eh_ato_fg_valido(texto: str) -> bool:
    assunto = limpar_assunto(texto)
    if re.search(r"secret[aá]rios?\s+de\s+estado", assunto or "", re.IGNORECASE):
        return bool(tem_codigo_vaga(texto) and eh_ato_relevante(texto))
    indicacao = extrair_indicacao(texto)
    if indicacao == "COLOCA" and re.search(r"disposi[cç]", texto or "", re.IGNORECASE):
        return bool(tem_codigo_vaga(texto))
    if indicacao == "RELOTA" and tem_codigo_vaga(texto):
        return True
    if indicacao == "RETIFICA" and (tem_codigo_vaga(texto) or eh_ato_relevante(texto)):
        return True
    if eh_assunto_fg_comissao(texto):
        return True
    if extrair_sem_efeito(texto) and eh_ato_relevante(texto):
        return True
    if not (tem_codigo_vaga(texto) and eh_ato_relevante(texto)):
        return False
    if assunto and ASSUNTO_IGNORAR_RE.search(assunto):
        return False
    return True


def extrair_assunto(texto: str) -> str:
    return ASSUNTO_PADRAO if eh_ato_fg_valido(texto) else ""


def extrair_indicacao(texto: str) -> str:
    sem = extrair_sem_efeito(texto)
    if sem:
        return sem
    t = texto.upper()
    for palavra in INDICACOES_ORDEM:
        if re.search(r"\b" + re.escape(palavra) + r"\b", t):
            return palavra
    return ""


def extrair_funcao_nivel(texto: str) -> str:
    return padronizar_codigo_funcao(texto)


def _limpa_trecho(texto: str) -> str:
    t = re.sub(r"\s+", " ", texto or "").strip(" ,.;")
    t = re.split(r"\s+a contar\b", t, maxsplit=1, flags=re.IGNORECASE)[0].strip(" ,.;")
    return t


def extrair_lotacao_relota(texto: str) -> str:
    dest = re.search(
        r"RELOTA\s+n[ao]\s+(.+?)(?:,\s*)?(?:a\s+)?(?:o\s+)?(?:Cargo|Fun[cç][aã]o)",
        texto or "",
        re.IGNORECASE | re.DOTALL,
    )
    orig = re.search(
        r"anteriormente lotad[oa]\s+n[ao]\s+(.+?)(?:\.|$)",
        texto or "",
        re.IGNORECASE,
    )
    destino = _limpa_trecho(dest.group(1)) if dest else ""
    origem = _limpa_trecho(orig.group(1)) if orig else ""
    if destino and origem:
        return f"{destino} (antes: {origem})"
    return destino or origem


def interpretar_coloca(texto: str) -> str:
    dest = re.search(
        r"à disposição do\(a\)\s+(.+?)(?:,|\s+a fim|\s+até|$)",
        texto or "",
        re.IGNORECASE | re.DOTALL,
    )
    if not dest:
        dest = re.search(
            r"à disposição d[aeo]\s+(.+?)(?:,|\s+a fim|\s+até|$)",
            texto or "",
            re.IGNORECASE | re.DOTALL,
        )
    fim = re.search(r"a fim de\s+(.+?)(?:,|\s+até|$)", texto or "", re.IGNORECASE | re.DOTALL)
    ate = re.search(r"até\s+(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})", texto or "", re.IGNORECASE)
    partes = []
    if dest:
        partes.append("à disposição de " + _limpa_trecho(dest.group(1)))
    if fim:
        partes.append(_limpa_trecho(fim.group(1)))
    if ate:
        partes.append("até " + ate.group(1))
    if re.search(r"ônus para o órgão de origem", texto or "", re.IGNORECASE):
        partes.append("com ônus para a origem")
    if re.search(r"ressarcimento", texto or "", re.IGNORECASE):
        partes.append("mediante ressarcimento")
    return "; ".join(partes)


def interpretar_retifica(texto: str) -> str:
    m = re.search(
        r"onde se lê[:\s]+(.+?)\s+leia-se[:\s]+(.+?)(?:\.|$)",
        texto or "",
        re.IGNORECASE | re.DOTALL,
    )
    if m:
        return (
            f"Onde se lê: {_limpa_trecho(m.group(1))} → "
            f"Leia-se: {_limpa_trecho(m.group(2))}"
        )
    m = re.search(r"para declarar que\s+(.+?)(?:\.|$)", texto or "", re.IGNORECASE)
    if m:
        return "Retifica: " + _limpa_trecho(m.group(1))
    m = re.search(r"RETIFICA\s+o ato.{0,220}", texto or "", re.IGNORECASE)
    if m:
        return _limpa_trecho(m.group(0))
    return "Retificação de ato"


def extrair_orgao_complemento(texto: str) -> str:
    padroes = [
        r"(?:com os |com |os )?encargos de\s+(.+?)(?:,\s*na vaga|\s*na vaga|\.|$)",
        r"com os encargos de\s+(.+?)(?:,|\.|$)",
        r"encargos de\s+(.+?)(?:,\s*na vaga|\s*na vaga)",
        r"na\s+Secretaria\s+(?:da|de)\s+([A-Za-záéíóúâêôãõç\s]+?)(?:,|\.| na vaga)",
    ]
    for p in padroes:
        m = re.search(p, texto, re.IGNORECASE | re.DOTALL)
        if m:
            return re.sub(r'\s+', ' ', m.group(1).strip()).strip(' ,.')
    return ""


def extrair_vaga_num(texto: str) -> str:
    padroes = [
        r"na\s+vaga\s*n[ºo°#]?\.?\s*([0-9]{6,})",
        r"vaga\s*n[ºo°#]?\.?\s*([0-9]{6,})",
        r"vaga\s*nº\.?\s*([0-9]{6,})",
        r"vaga\s*n°\.?\s*([0-9]{6,})",
        r"vaga\s*n\.º\s*([0-9]{6,})",
        r"vaga\s*n\.°\s*([0-9]{6,})",
        r"na\s+vaga\s*([0-9]{6,})",
        r"vaga\s*nº\s*([0-9]{6,})",
        r"vaga\s*n°\s*([0-9]{6,})",
        r"vaga\s*n[ºo°]\s*([0-9]{6,})",
    ]
    for p in padroes:
        m = re.search(p, texto, re.IGNORECASE)
        if m:
            return m.group(1)
    return ""


def extrair_vaga_anterior(texto: str) -> str:
    padroes = [
        r"anteriormente ocupada por\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-záéíóúâêôãõç\s]+?)(?:,|\.| identificação|$)",
        r"deixada por\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-záéíóúâêôãõç\s]+?)(?:,|\.| identificação|$)",
        r"ocupada por\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-záéíóúâêôãõç\s]+?)(?:,|\.| identificação|$)",
        r"deixada\s+por\s+([A-ZÁÉÍÓÚÂÊÔÃÕÇ][A-Za-záéíóúâêôãõç\s]+?)(?:,|\.|$)",
    ]
    for p in padroes:
        m = re.search(p, texto, re.IGNORECASE)
        if m:
            nome = m.group(1).strip()
            nome = re.split(r',|\s+identifica', nome, flags=re.I)[0].strip()
            return re.sub(r'\s+', ' ', nome).strip()
    return ""


def normalizar_data_br(texto: str) -> str:
    m = re.match(r"(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})$", (texto or "").strip())
    if not m:
        return ""
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    if len(y) == 2:
        y = "20" + y
    return f"{d:02d}/{mo:02d}/{y}"


def extrair_a_contar_de(texto: str) -> str:
    m = re.search(
        r"a\s+contar\s+de\s+(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})",
        texto or "",
        re.IGNORECASE,
    )
    if m:
        return normalizar_data_br(m.group(1))
    return ""


def extrair_proa(texto: str) -> str:
    m = re.search(r"Processo:\s*([0-9/.\-]+)", texto, re.IGNORECASE)
    return m.group(1).strip() if m else ""


def eh_primeiro_provimento(texto: str) -> bool:
    return bool(re.search(r"primeiro\s+provimento", texto or "", re.IGNORECASE))


def eh_celula_meta(texto: str) -> bool:
    t = (texto or "").strip()
    if not t or t in {"-", "–"}:
        return True
    tl = t.lower()
    if tl in {"a contar de", "a contar", "identificação", "identificacao", "protocolo"}:
        return True
    if re.match(r"a\s+contar\s+de$", tl):
        return True
    if re.match(r"\d{1,2}/\d{1,2}/\d{4}$", t):
        return True
    if ID_FUNCIONAL_RE.fullmatch(t):
        return True
    return False


def limpar_campo_resultado(texto: str) -> str:
    if not texto:
        return ""
    t = str(texto)
    t = re.sub(r"Protocolo:\s*\d+", "", t, flags=re.IGNORECASE)
    t = re.sub(r"A\s+contar\s+de\s+\d{1,2}/\d{1,2}/\d{4}", "", t, flags=re.IGNORECASE)
    t = re.sub(r"Identifica[cç][aã]o(?:\s+Funcional)?(?:/V[ií]nculo)?\s*:?\s*\S*", "", t, flags=re.IGNORECASE)
    t = re.sub(r"\b\d{5,8}\s*/\s*\d{1,2}\b", "", t)
    t = re.split(r"\s*[-–,]?\s*Id\.?\s*Func\.?", t, maxsplit=1, flags=re.IGNORECASE)[0]
    t = re.split(r"\s*[-–,]?\s*ID\.?\s*n[ºo°]?", t, maxsplit=1, flags=re.IGNORECASE)[0]
    return re.sub(r"\s+", " ", t).strip(" -–,.")


def vaga_anterior_resultado(valor: str, texto_ref: str, indicacao: str) -> str:
    if str(indicacao or "").upper() in INDICACOES_DESOCUPA:
        return ""
    blob = f"{valor or ''} {texto_ref or ''}"
    if eh_primeiro_provimento(blob):
        return "Primeiro Provimento"
    t = limpar_campo_resultado(valor or "")
    if eh_celula_meta(t) or re.search(r"a\s+contar", t, re.IGNORECASE):
        return ""
    if re.search(r"assessor|diretor|chefe|coordenador", t, re.I) and len(t) < 50:
        return ""
    return t


def limpar_celula(valor) -> str:
    if valor is None:
        return ""
    return re.sub(r"\s+", " ", str(valor)).strip()


def eh_cabecalho_quadro(celulas: List[str]) -> bool:
    junto = " ".join(celulas).lower()
    if "assunto:" in junto or "protocolo:" in junto:
        return False
    if len(junto) > 90:
        return False
    if not re.search(r"\bnome\b", junto):
        return False
    return any(
        re.search(rf"\b{p}\b", junto)
        for p in (
            "código", "codigo", "vaga", "encargo", "titular",
            "identifica", "lotação", "lotacao", "contar",
        )
    )


def eh_nome_pessoa(nome: str) -> bool:
    """Rejeita resto de frase/quebra de página (ex.: 'Secretaria da …, na vaga n.º …')."""
    n = re.sub(r"\s+", " ", (nome or "")).strip()
    if len(n) < 5 or len(n) > 80:
        return False
    if n in {"—", "-"}:
        return False
    if n.lower() in {"nome", "município", "municipio", "a contar de", "identificação", "identificacao"}:
        return False
    if eh_celula_meta(n):
        return False
    nu = n.upper()
    if any(p in nu for p in ("GOVERNADOR", "PROTOCOLO", "ASSUNTO:", "PROCESSO:")):
        return False
    if re.search(r"\bvaga\s*n", n, re.IGNORECASE):
        return False
    if re.search(r"n\.?\s*[ºo°]\s*\d", n, re.IGNORECASE):
        return False
    if re.match(
        r"^(SECRETARIA|FUNDA[CÇ][AÃ]O|INSTITUTO|GABINETE|DEPARTAMENTO|COORDENADORIA|POL[IÍ]CIA|CASA CIVIL|DI[AÁ]RIO OFICIAL)\b",
        nu,
    ):
        return False
    palavras = [p for p in n.split() if re.search(r"[A-Za-zÁ-ú]", p)]
    return len(palavras) >= 2


def eh_linha_quadro_fg(celulas: List[str], exigir_codigo: bool = True) -> bool:
    if not celulas:
        return False
    junto = " ".join(celulas)
    if not VAGA_QUADRO_RE.search(junto):
        return False
    if exigir_codigo and not CODIGO_QUADRO_RE.search(junto):
        return False
    return eh_nome_pessoa(celulas[0])


def codigo_do_preambulo(texto: str) -> str:
    return padronizar_codigo_funcao(texto)


def limpar_titular(texto: str) -> str:
    if not texto:
        return ""
    texto = re.split(r"\s*[-–,]?\s*Id\.?\s*Func\.?", texto, flags=re.IGNORECASE)[0]
    texto = re.split(r"\s*[-–,]?\s*ID\.?\s*n[ºo°]?", texto, flags=re.IGNORECASE)[0]
    texto = re.sub(r"\s+", " ", texto).strip(" -–,.")
    if re.search(r"assessor|diretor|chefe|coordenador|secretário|secretario", texto, re.I):
        return ""
    return texto


def parse_linha_quadro(celulas: List[str], codigo_padrao: str = "") -> Optional[Dict]:
    celulas = [limpar_celula(c) for c in celulas]
    junto = " ".join(celulas)
    m_cod = CODIGO_QUADRO_RE.search(junto)
    if m_cod:
        codigo = f"{m_cod.group(1).upper()}-{int(m_cod.group(2)):02d}"
    elif codigo_padrao:
        codigo = codigo_padrao
    else:
        return None
    vagas = VAGA_QUADRO_RE.findall(junto)
    vaga = vagas[-1] if vagas else ""
    nome = celulas[0]
    a_contar = ""
    for c in celulas:
        if re.match(r"\d{1,2}/\d{1,2}/\d{4}$", c or ""):
            a_contar = normalizar_data_br(c)
            break

    titular = ""
    for candidato in reversed(celulas[1:]):
        if not candidato or candidato == nome or eh_celula_meta(candidato):
            continue
        if eh_primeiro_provimento(candidato):
            titular = "Primeiro Provimento"
            break
        if VAGA_QUADRO_RE.fullmatch(candidato) or CODIGO_QUADRO_RE.fullmatch(candidato.replace(" ", "")):
            continue
        if (
            re.search(r"Id\.?\s*Func|ID\.?\s*n", candidato, re.IGNORECASE)
            or (not candidato.isdigit() and sum(ch.isalpha() for ch in candidato) >= 6)
        ):
            titular = limpar_titular(candidato)
            if titular:
                break

    encargo = ""
    meio = celulas[1:-1] if titular else celulas[1:]
    for c in meio:
        if not c:
            continue
        if CODIGO_QUADRO_RE.fullmatch(re.sub(r"\s+", "", c).replace("–", "-")):
            continue
        if VAGA_QUADRO_RE.fullmatch(c):
            continue
        if ID_FUNCIONAL_RE.fullmatch(c) or eh_celula_meta(c) or eh_primeiro_provimento(c):
            continue
        if CODIGO_QUADRO_RE.search(c) and VAGA_QUADRO_RE.search(c) and len(c) < 45:
            continue
        encargo = c
        break

    orgao = ""
    funcao_desc = encargo
    m_org = re.search(r"\s*[-–]\s*(Na\s+.+)$", encargo, re.IGNORECASE)
    if m_org:
        orgao = m_org.group(1).strip()
        funcao_desc = encargo[: m_org.start()].strip()

    codigo_abrev = padronizar_codigo_funcao(codigo)
    funcao = codigo_abrev
    if funcao_desc and padronizar_codigo_funcao(funcao_desc) == funcao_desc:
        funcao = f"{codigo_abrev} – {funcao_desc}".strip(" –") if codigo_abrev else funcao_desc
    return {
        "Nome": re.sub(r"\s+", " ", nome).strip(),
        "Vaga n°": vaga,
        "Função": funcao,
        "orgao_encargo": orgao,
        "Vaga Anterior": titular,
        "A contar de": a_contar,
    }


def ultimo_match(padrao: str, texto: str) -> Optional[re.Match]:
    achados = list(re.finditer(padrao, texto, re.IGNORECASE))
    return achados[-1] if achados else None


def extrair_indicacao_ultima(texto: str) -> str:
    if not texto:
        return ""
    candidatos = []
    for m in SEM_EFEITO_RE.finditer(texto):
        candidatos.append((m.start(), f"{m.group(1).upper()} SEM EFEITO"))
    t = texto.upper()
    for palavra in INDICACOES_ORDEM:
        for m in re.finditer(r"\b" + re.escape(palavra) + r"\b", t):
            candidatos.append((m.start(), palavra))
    if not candidatos:
        return ""
    return max(candidatos, key=lambda item: item[0])[1]


def extrair_secretaria_contexto(texto: str) -> str:
    m = ultimo_match(
        r"\b((?:na|no)\s+(?:Secretaria|Corpo|Instituto|Fundação|Fundacao|Polícia|Policia|Brigada)[^,:\.\n]{3,70})",
        texto,
    )
    if not m:
        return ""
    return re.sub(r"\s+", " ", m.group(1)).strip(" :.")


def atualizar_contexto_quadro(texto: str, ctx: Dict) -> Dict:
    if not texto:
        return ctx
    if re.search(r"Protocolo:\s*\d+", texto):
        ctx["sem_efeito"] = extrair_sem_efeito(texto)
        if "Assunto:" in texto or ctx["sem_efeito"]:
            ctx["assunto_ok"] = eh_ato_fg_valido(texto)
        else:
            ctx["assunto_ok"] = False
    elif "Assunto:" in texto:
        ctx["assunto_ok"] = eh_ato_fg_valido(texto)
        ctx["sem_efeito"] = extrair_sem_efeito(texto)
    else:
        sem = extrair_sem_efeito(texto)
        if sem:
            ctx["sem_efeito"] = sem
            if eh_ato_relevante(texto):
                ctx["assunto_ok"] = True
    m_proa = ultimo_match(r"Processo:\s*([0-9/.\-]+)", texto)
    if m_proa:
        ctx["Proa"] = m_proa.group(1).strip()
    m_prot = ultimo_match(r"Protocolo:\s*(\d+)", texto)
    if m_prot:
        ctx["protocolo"] = m_prot.group(1)
    indicacao = extrair_indicacao_ultima(texto)
    if indicacao:
        ctx["Indicação"] = indicacao
    codigo = codigo_do_preambulo(texto)
    if codigo:
        ctx["codigo"] = codigo
    secretaria = extrair_secretaria_contexto(texto)
    if secretaria:
        ctx["secretaria"] = secretaria
    return ctx


def iter_tabelas(pagina: Dict):
    for item in pagina.get("tabelas") or []:
        if isinstance(item, dict):
            yield item.get("preamble") or "", item.get("rows") or []
        else:
            yield "", item or []


def texto_antes_do_primeiro_quadro(texto: str, tabelas: List) -> str:
    primeiro_nome = ""
    for item in tabelas or []:
        rows = item.get("rows") if isinstance(item, dict) else item
        for bruto in rows or []:
            celulas = [limpar_celula(c) for c in (bruto or [])]
            if eh_linha_quadro_fg(celulas, exigir_codigo=False):
                primeiro_nome = celulas[0]
                break
        if primeiro_nome:
            break
    if not primeiro_nome:
        return texto
    idx = texto.upper().find(primeiro_nome.upper())
    if idx < 0:
        return texto
    return texto[:idx]


def secao_da_pagina(pagina_num: int, secoes_sumario: List[Tuple[str, int]]) -> str:
    secao = "Outras Seções"
    for nome_secao, pagina_inicio in reversed(secoes_sumario):
        if pagina_num >= pagina_inicio:
            return nome_secao
    return secao


def aplicar_regra_exoneracao(reg: Dict) -> Dict:
    if str(reg.get("Indicação", "")).strip().upper() in INDICACOES_DESOCUPA:
        reg["Vaga Anterior"] = ""
    return reg


def chave_data(data: str) -> Tuple:
    texto = str(data or "").strip()
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", texto)
    if not m:
        return (9999, 99, 99, texto)
    return (int(m.group(3)), int(m.group(2)), int(m.group(1)), "")


def vaga_normalizada(reg: Dict) -> str:
    return re.sub(r"\D", "", str(reg.get("Vaga n°") or ""))


def cruzar_movimentacoes_por_vaga(registros: List[Dict]) -> List[Dict]:
    """Liga saída (EXONERA/DISPENSA/CESSA) à entrada (NOMEIA/DESIGNA/ATRIBUI) pelo código da vaga.

    No mesmo DOE a ordem no PDF não importa. Também cruza datas diferentes.
    Se a entrada não tiver Vaga Anterior, preenche com quem saiu.
    """
    por_vaga: Dict[str, List[Dict]] = defaultdict(list)
    for reg in registros:
        vaga = vaga_normalizada(reg)
        if len(vaga) >= 9:
            por_vaga[vaga].append(reg)

    pares: List[Dict] = []
    for vaga, atos in sorted(por_vaga.items()):
        saidas = [r for r in atos if str(r.get("Indicação", "")).upper() in INDICACOES_DESOCUPA]
        entradas = [r for r in atos if str(r.get("Indicação", "")).upper() in INDICACOES_ENTRADA_VAGA]
        saidas.sort(key=lambda r: chave_data(r.get("Data DOE", "")))
        entradas.sort(key=lambda r: chave_data(r.get("Data DOE", "")))
        usados_s: set = set()
        usados_e: set = set()

        def monta_par(saida: Dict, entrada: Dict) -> Dict:
            if not str(entrada.get("Vaga Anterior") or "").strip():
                entrada["Vaga Anterior"] = saida.get("Nome") or ""
            ds = str(saida.get("Data DOE") or "")
            de = str(entrada.get("Data DOE") or "")
            ac_s = str(saida.get("A contar de") or "")
            ac_e = str(entrada.get("A contar de") or "")
            if ac_s and ac_e and ac_s == ac_e:
                sit = "Pareado pela data a contar de"
            elif ds == de:
                sit = "Pareado no mesmo DOE"
            else:
                sit = "Pareado em datas diferentes"
            if de and ds and chave_data(de) < chave_data(ds):
                sit += " (nomeação publicada antes da exoneração)"
            return {
                "Vaga n°": vaga,
                "Quem saiu": saida.get("Nome", ""),
                "Indicação saída": saida.get("Indicação", ""),
                "Data saída": ds,
                "A contar saída": ac_s,
                "Quem entrou": entrada.get("Nome", ""),
                "Indicação entrada": entrada.get("Indicação", ""),
                "Data entrada": de,
                "A contar entrada": ac_e,
                "Vaga Anterior": entrada.get("Vaga Anterior", ""),
                "Situação": sit,
            }

        acontar_vals = sorted(
            {str(r.get("A contar de") or "") for r in saidas + entradas if r.get("A contar de")},
            key=chave_data,
        )
        for ac in acontar_vals:
            s_ac = [r for r in saidas if str(r.get("A contar de") or "") == ac and id(r) not in usados_s]
            e_ac = [r for r in entradas if str(r.get("A contar de") or "") == ac and id(r) not in usados_e]
            for saida, entrada in zip(s_ac, e_ac):
                usados_s.add(id(saida))
                usados_e.add(id(entrada))
                pares.append(monta_par(saida, entrada))

        datas = sorted(
            {str(r.get("Data DOE") or "") for r in saidas + entradas},
            key=chave_data,
        )
        for data in datas:
            s_dia = [r for r in saidas if str(r.get("Data DOE") or "") == data and id(r) not in usados_s]
            e_dia = [r for r in entradas if str(r.get("Data DOE") or "") == data and id(r) not in usados_e]
            for saida, entrada in zip(s_dia, e_dia):
                usados_s.add(id(saida))
                usados_e.add(id(entrada))
                pares.append(monta_par(saida, entrada))

        s_rest = [r for r in saidas if id(r) not in usados_s]
        e_rest = [r for r in entradas if id(r) not in usados_e]
        for saida in s_rest:
            if not e_rest:
                pares.append({
                    "Vaga n°": vaga,
                    "Quem saiu": saida.get("Nome", ""),
                    "Indicação saída": saida.get("Indicação", ""),
                    "Data saída": saida.get("Data DOE", ""),
                    "A contar saída": saida.get("A contar de", ""),
                    "Quem entrou": "",
                    "Indicação entrada": "",
                    "Data entrada": "",
                    "A contar entrada": "",
                    "Vaga Anterior": "",
                    "Situação": "Só saída — conferir nomeação noutro DOE",
                })
                continue
            ds = chave_data(saida.get("A contar de") or saida.get("Data DOE", ""))

            def chave_entrada(entrada: Dict):
                de = chave_data(entrada.get("A contar de") or entrada.get("Data DOE", ""))
                dist = abs((de[0] - ds[0]) * 372 + (de[1] - ds[1]) * 31 + (de[2] - ds[2]))
                return dist

            entrada = min(e_rest, key=chave_entrada)
            e_rest.remove(entrada)
            pares.append(monta_par(saida, entrada))
        for entrada in e_rest:
            pares.append({
                "Vaga n°": vaga,
                "Quem saiu": "",
                "Indicação saída": "",
                "Data saída": "",
                "A contar saída": "",
                "Quem entrou": entrada.get("Nome", ""),
                "Indicação entrada": entrada.get("Indicação", ""),
                "Data entrada": entrada.get("Data DOE", ""),
                "A contar entrada": entrada.get("A contar de", ""),
                "Vaga Anterior": entrada.get("Vaga Anterior", ""),
                "Situação": "Só entrada — conferir exoneração noutro DOE",
            })
    return pares


def chave_registro(reg: Dict) -> Tuple[str, str, str]:
    nome = re.sub(r"\s+", " ", reg.get("Nome", "")).strip().upper()
    vaga = re.sub(r"\D", "", reg.get("Vaga n°", ""))
    ind = str(reg.get("Indicação", "")).strip().upper()
    return nome, vaga, ind


def extrair_registros_quadros(
    paginas: List[Dict],
    data_doe: str,
    secoes_sumario: List[Tuple[str, int]],
) -> List[Dict]:
    ctx: Dict = {
        "Indicação": "",
        "Proa": "",
        "protocolo": "",
        "secretaria": "",
        "assunto_ok": False,
        "sem_efeito": "",
        "codigo": "",
    }
    registros = []

    for pagina in paginas:
        pagina_num = pagina["pagina"]
        secao = secao_da_pagina(pagina_num, secoes_sumario)
        prefixo = texto_antes_do_primeiro_quadro(pagina.get("texto", ""), pagina.get("tabelas") or [])
        atualizar_contexto_quadro(prefixo, ctx)

        for preamble, tabela in iter_tabelas(pagina):
            if preamble:
                atualizar_contexto_quadro(preamble, ctx)
            if not tabela:
                continue
            exigir_codigo = not bool(ctx.get("Indicação") and ctx.get("assunto_ok"))
            for bruto in tabela:
                celulas = [limpar_celula(c) for c in (bruto or [])]
                if not any(celulas):
                    continue
                junto = " ".join(celulas)
                if eh_cabecalho_quadro(celulas):
                    continue
                if "Protocolo:" in junto or "Assunto:" in junto or extrair_indicacao(junto):
                    atualizar_contexto_quadro(junto, ctx)
                    exigir_codigo = not bool(ctx.get("Indicação") and ctx.get("assunto_ok"))
                    if not eh_linha_quadro_fg(celulas, exigir_codigo=exigir_codigo):
                        continue
                if not eh_linha_quadro_fg(celulas, exigir_codigo=exigir_codigo):
                    continue

                if not ctx.get("assunto_ok"):
                    continue

                parsed = parse_linha_quadro(celulas, codigo_padrao=ctx.get("codigo", ""))
                if not parsed:
                    continue

                orgao = parsed.pop("orgao_encargo", "") or ctx.get("secretaria", "")
                indicacao = ctx.get("Indicação", "")
                texto_ref = f"{preamble} {junto}"
                registros.append(aplicar_regra_exoneracao({
                    "Nome": limpar_campo_resultado(parsed["Nome"]),
                    "Assunto": ASSUNTO_PADRAO,
                    "Indicação nova": "Sim" if indicacao in INDICACOES_NOVA else "",
                    "Data DOE": data_doe,
                    "A contar de": parsed.get("A contar de") or extrair_a_contar_de(texto_ref),
                    "Indicação": indicacao,
                    "Partido dono da vaga": "",
                    "Deputado": "",
                    "Vaga Anterior": vaga_anterior_resultado(
                        parsed["Vaga Anterior"], texto_ref, indicacao
                    ),
                    "Proa": ctx.get("Proa", ""),
                    "Vaga n°": parsed["Vaga n°"],
                    "Função": padronizar_codigo_funcao(parsed["Função"]),
                    "Órgão / Secretaria": limpar_campo_resultado(orgao),
                    "_secao": secao,
                }))
        atualizar_contexto_quadro(pagina.get("texto_apos_tabelas") or "", ctx)
    return registros


def eh_ato_relevante(texto: str) -> bool:
    texto_upper = texto.upper()
    indicadores = [
        "CARGO COMISSIONADO", "FUNÇÃO GRATIFICADA", "FUNCAO GRATIFICADA",
        "CARGOS EM COMISSÃO", "CARGOS EM COMISSAO",
        "FUNÇÕES GRATIFICADAS", "FUNCOES GRATIFICADAS",
        "CCS", "CCT", "FGT", "FGS",
        "CARGO COMISS.", "FUNÇÃO GRATIF.", "FUNCAO GRATIF.",
        "GRATIFICAÇÃO DE CONFIANÇA", "GRATIFICACAO DE CONFIANCA",
    ]
    return any(ind in texto_upper for ind in indicadores)


def processar_bloco(texto: str, data_doe: str, secao_atual: str) -> Optional[Dict]:
    if "Protocolo:" not in texto and "Assunto:" not in texto:
        return None
    if not eh_ato_fg_valido(texto):
        return None
    if eh_assunto_fg_comissao(texto) and not eh_ato_relevante(texto):
        return None

    indicacao = extrair_indicacao(texto)
    nome = limpar_campo_resultado(extrair_nome(texto))
    if not nome:
        if indicacao == "RELOTA":
            nome = "—"
        else:
            return None
    if nome != "—" and not eh_nome_pessoa(nome):
        return None

    orgao = extrair_orgao_complemento(texto)
    ant = extrair_vaga_anterior(texto)
    assunto = ASSUNTO_PADRAO
    if indicacao == "RELOTA":
        orgao = extrair_lotacao_relota(texto) or orgao
        ant = ""
    elif indicacao == "COLOCA":
        orgao = interpretar_coloca(texto) or orgao
        ant = ""
        if re.search(r"disposi", limpar_assunto(texto) or "", re.IGNORECASE):
            assunto = "Disposição no Estado"
    elif indicacao == "RETIFICA":
        orgao = interpretar_retifica(texto) or orgao

    return {
        "Nome": nome,
        "Assunto": assunto,
        "Indicação nova": "Sim" if indicacao in INDICACOES_NOVA else "",
        "Data DOE": data_doe,
        "A contar de": extrair_a_contar_de(texto),
        "Indicação": indicacao,
        "Partido dono da vaga": "",
        "Deputado": "",
        "Vaga Anterior": vaga_anterior_resultado(ant, texto, indicacao),
        "Proa": extrair_proa(texto),
        "Vaga n°": extrair_vaga_num(texto),
        "Função": padronizar_codigo_funcao(extrair_funcao_nivel(texto)),
        "Órgão / Secretaria": limpar_campo_resultado(orgao),
        "_secao": secao_atual
    }


def extrair_paginas(caminho_pdf: str) -> List[Dict]:
    resultado = []
    try:
        with pdfplumber.open(caminho_pdf) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                itens = []
                encontrados = page.find_tables() or []
                prev_y = 0
                for tabela in encontrados:
                    bbox = tabela.bbox
                    y0 = bbox[1]
                    preamble = ""
                    if y0 > prev_y + 1:
                        try:
                            preamble = page.crop((0, prev_y, page.width, y0)).extract_text() or ""
                        except Exception:
                            preamble = ""
                    itens.append({
                        "bbox": bbox,
                        "rows": tabela.extract() or [],
                        "preamble": preamble,
                    })
                    prev_y = bbox[3]
                texto_apos = ""
                if itens:
                    last_y = itens[-1]["bbox"][3]
                    if last_y < page.height - 2:
                        try:
                            texto_apos = page.crop((0, last_y, page.width, page.height)).extract_text() or ""
                        except Exception:
                            texto_apos = ""
                resultado.append({
                    "pagina": i,
                    "texto": page.extract_text() or "",
                    "tabelas": itens,
                    "texto_apos_tabelas": texto_apos,
                })
    except Exception as e:
        print(f"  [ERRO] Falha ao extrair {caminho_pdf}: {e}")
        return []
    return resultado


def processar_pdf(caminho_pdf: Path) -> Tuple[List[Dict], List[str]]:
    print(f"\nProcessando: {caminho_pdf.name}")

    paginas = extrair_paginas(str(caminho_pdf))
    if not paginas:
        return [], []

    secoes_sumario = extrair_sumario(paginas)
    secoes_ordenadas = [s[0] for s in secoes_sumario]

    texto_completo = "\n".join([p["texto"] for p in paginas])
    data_doe = extrair_data_doe(texto_completo)

    protocolo_pagina = {}
    for pagina in paginas:
        for match in re.finditer(r"Protocolo:\s*(\d+)", pagina["texto"]):
            protocolo = match.group(1)
            if protocolo not in protocolo_pagina:
                protocolo_pagina[protocolo] = pagina["pagina"]

    registros = []
    for parte in re.split(r"(?=Protocolo:\s*\d+)|(?=Assunto:\s)", texto_completo):
        if "Protocolo:" not in parte and "Assunto:" not in parte:
            continue

        match_prot = re.search(r"Protocolo:\s*(\d+)", parte)
        if match_prot:
            pagina_do_ato = protocolo_pagina.get(match_prot.group(1), 1)
        else:
            trecho = re.sub(r"\s+", " ", parte)[:80]
            pagina_do_ato = 1
            for pagina in paginas:
                if trecho and trecho in re.sub(r"\s+", " ", pagina["texto"]):
                    pagina_do_ato = pagina["pagina"]
                    break
        secao_atual = secao_da_pagina(pagina_do_ato, secoes_sumario)

        reg = processar_bloco(parte, data_doe, secao_atual)
        if reg:
            registros.append(reg)

    vistos = {chave_registro(r) for r in registros}
    quadros = extrair_registros_quadros(paginas, data_doe, secoes_sumario)
    novos_quadros = 0
    for reg in quadros:
        chave = chave_registro(reg)
        if chave in vistos and chave[0]:
            continue
        vistos.add(chave)
        registros.append(reg)
        novos_quadros += 1

    print(f"  → {len(registros)} registros extraídos ({len(registros) - novos_quadros} atos + {novos_quadros} quadros)")
    return registros, secoes_ordenadas


def criar_planilha(registros: List[Dict], secoes_ordenadas: List[str], caminho: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "FG e Cargos em Comissão"

    AZUL = "1A365D"
    header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=AZUL)
    cell_font = Font(name="Arial", size=10)
    secao_font = Font(name="Arial", bold=True, size=11, color="FFFFFF")
    secao_fill = PatternFill("solid", fgColor="2E4053")
    data_font = Font(name="Arial", bold=True, size=12, color="FFFFFF")
    data_fill = PatternFill("solid", fgColor="1A5276")
    thin = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0")
    )

    ws.merge_cells("A1:M1")
    ws["A1"] = "FUNÇÕES GRATIFICADAS E CARGOS EM COMISSÃO - RELATÓRIO"
    ws["A1"].font = Font(name="Arial", bold=True, size=14, color=AZUL)

    ws.merge_cells("A2:M2")
    ws["A2"] = f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  Total de registros: {len(registros)}"

    for col, h in enumerate(COLUNAS, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[4].height = 32

    for i, w in enumerate([28, 35, 12, 12, 12, 12, 16, 16, 22, 16, 14, 18, 40], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    por_data = defaultdict(lambda: defaultdict(list))
    for reg in registros:
        data = reg.get("Data DOE", "Sem data")
        secao = reg.get("_secao", "Outras Seções")
        por_data[data][secao].append(reg)

    datas_ordenadas = sorted(por_data.keys(), key=chave_data)
    linha_atual = 5

    for data in datas_ordenadas:
        ws.merge_cells(f"A{linha_atual}:M{linha_atual}")
        cell = ws.cell(row=linha_atual, column=1, value=f"DATA DO DOE: {data}")
        cell.font = data_font
        cell.fill = data_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[linha_atual].height = 24
        linha_atual += 1

        secoes_extra = [s for s in por_data[data] if s not in secoes_ordenadas]
        for secao in list(secoes_ordenadas) + secoes_extra:
            atos = por_data[data].get(secao, [])
            if not atos:
                continue

            ws.merge_cells(f"A{linha_atual}:M{linha_atual}")
            cell = ws.cell(row=linha_atual, column=1, value=secao.upper())
            cell.font = secao_font
            cell.fill = secao_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[linha_atual].height = 22
            linha_atual += 1

            for reg in atos:
                for c_idx, chave in enumerate(COLUNAS, 1):
                    cell = ws.cell(row=linha_atual, column=c_idx, value=reg.get(chave, ""))
                    cell.font = cell_font
                    cell.border = thin
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if linha_atual % 2 == 0:
                        cell.fill = PatternFill("solid", fgColor="F7FAFC")
                linha_atual += 1

    ws.freeze_panes = "A5"

    pares = cruzar_movimentacoes_por_vaga(registros)
    criar_aba_movimentacoes(wb, pares)

    wb.save(caminho)
    print(f"\n✓ Arquivo salvo: {caminho}")
    n_par = sum(1 for p in pares if p["Situação"].startswith("Pareado"))
    n_rev = len(pares) - n_par
    print(f"  Movimentações por vaga: {n_par} pareadas, {n_rev} para conferir")


def criar_aba_movimentacoes(wb: Workbook, pares: List[Dict]):
    ws = wb.create_sheet("Movimentações por vaga")
    cols = [
        "Vaga n°", "Quem saiu", "Indicação saída", "Data saída", "A contar saída",
        "Quem entrou", "Indicação entrada", "Data entrada", "A contar entrada",
        "Vaga Anterior", "Situação",
    ]
    header_font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1A365D")
    cell_font = Font(name="Arial", size=10)
    thin = Border(
        left=Side(style="thin", color="CBD5E0"),
        right=Side(style="thin", color="CBD5E0"),
        top=Side(style="thin", color="CBD5E0"),
        bottom=Side(style="thin", color="CBD5E0"),
    )
    fill_ok = PatternFill("solid", fgColor="E6F4EA")
    fill_rev = PatternFill("solid", fgColor="FFF3CD")

    ws.merge_cells("A1:K1")
    ws["A1"] = (
        "Cruzamento por código de vaga — EXONERA/DISPENSA/CESSA × NOMEIA/DESIGNA/ATRIBUI. "
        "A contar de é a data de efeito; a nomeação pode ser publicada antes da exoneração."
    )
    ws["A1"].font = Font(name="Arial", bold=True, size=12, color="1A365D")
    ws.merge_cells("A2:K2")
    ws["A2"] = (
        "Cruza pelo número da vaga. Ordem no DOE não importa. "
        "Prioriza a mesma data a contar de; depois o mesmo DOE; depois datas diferentes."
    )

    for col, h in enumerate(cols, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 28
    for i, w in enumerate([16, 28, 14, 12, 12, 28, 14, 12, 12, 28, 42], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    linha = 5
    for par in pares:
        situacao = par.get("Situação", "")
        fundo = fill_ok if situacao.startswith("Pareado") else fill_rev
        for c_idx, chave in enumerate(cols, 1):
            cell = ws.cell(row=linha, column=c_idx, value=par.get(chave, ""))
            cell.font = cell_font
            cell.border = thin
            cell.fill = fundo
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        linha += 1
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:K{max(4, linha - 1)}"


def main():
    os.chdir(PASTA_AGENTE)

    parser = argparse.ArgumentParser(description="Agente FG / Cargos em Comissão")
    parser.add_argument(
        "--pdf",
        nargs="+",
        metavar="ARQUIVO",
        help="Um ou mais PDFs para processar (sem menu interativo)",
    )
    args = parser.parse_args()

    print("=" * 60)
    print("  AGENTE FG / CARGOS EM COMISSÃO")
    print("=" * 60)

    if args.pdf:
        selecionados = resolver_pdfs_cli(args.pdf)
    else:
        selecionados = selecionar_pdfs(listar_pdfs())

    print(f"\nArquivos selecionados ({len(selecionados)}):")
    for pdf in selecionados:
        print(f"  - {pdf.name}")

    todos_registros = []
    todas_secoes = []

    for pdf in selecionados:
        registros, secoes = processar_pdf(pdf)
        todos_registros.extend(registros)
        for s in secoes:
            if s not in todas_secoes:
                todas_secoes.append(s)

    unicos = []
    vistos = set()
    for reg in todos_registros:
        chave = chave_registro(reg)
        if chave in vistos and (chave[0] or chave[1]):
            continue
        vistos.add(chave)
        unicos.append(reg)
    todos_registros = unicos

    print(f"\n{'='*60}")
    print(f"Total geral de registros extraídos: {len(todos_registros)}")

    if todos_registros:
        nome_excel = PASTA_AGENTE / f"FG_Cargos_Comissao_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
        criar_planilha(todos_registros, todas_secoes, str(nome_excel))
    else:
        print("Nenhum registro encontrado nos arquivos selecionados.")

    print("\nProcessamento concluído.")


if __name__ == "__main__":
    main()